import openpyxl

wb = openpyxl.load_workbook('Sidewalk_Mosaic_Nonprofit_Financial_Model.xlsx', data_only=True)
for name in wb.sheetnames:
    print(f'Sheet: {name}')
    ws = wb[name]
    for row in ws.iter_rows():
        values = [c.value for c in row]
        if any(v is not None for v in values):
            print(values)
    print()
